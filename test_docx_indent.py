import docx
from docx.shared import Inches

doc = docx.Document()
p1 = doc.add_paragraph("No indent")
p2 = doc.add_paragraph("0.5 inch indent")
p2.paragraph_format.left_indent = Inches(0.5)
p3 = doc.add_paragraph("1 inch indent")
p3.paragraph_format.left_indent = Inches(1.0)
doc.save("test_indent.docx")

doc_read = docx.Document("test_indent.docx")
for p in doc_read.paragraphs:
    print(p.text, p.paragraph_format.left_indent)
